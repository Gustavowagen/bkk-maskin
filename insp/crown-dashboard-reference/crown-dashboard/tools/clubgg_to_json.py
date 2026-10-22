"""
Convert a ClubGG weekly club export (.xlsx) into the JSON datasets the dashboard reads.

Usage:
    python tools/clubgg_to_json.py path/to/export.xlsx [out_dir]

Writes <out_dir>/sessions.json, player_sessions.json, members.json, mtt.json,
mtt_results.json and data.js (window.CK_DATA = {...}). out_dir defaults to ./data.
Requires: pip install openpyxl

Sheets read: Club Overview, Member Statistics, Club Member Balance,
Ring Game Detail, MTT Statistics, MTT Detail. All times stay in the export's
own clock (UTC-7); the page shifts them to Oslo time for display.
"""
import json, os, re, sys
from datetime import datetime
import openpyxl

IDRE = re.compile(r'^\d{4}-\d{4}$')
num = lambda v: float(v) if v not in (None, '', '-') else 0.0
dash_ = lambda v: None if v in (None, '-', '') else v


def read_export(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ov = wb['Club Overview']
    meta = {
        'club': ov['A1'].value.split(':', 1)[1].strip(),
        'clubId': ov['A2'].value.split(':', 1)[1].strip(),
        'period': ov['A3'].value.split(':', 1)[1].strip(),
    }

    # Club Overview: one row per member, week totals.
    # cols: 1/2 super agent id/nick, 3/4 agent id/nick, 6 role, 7/8 member id/nick,
    #       9 games, 10 hands, 11 fee, 15 BBJ contribution, 19 P&L
    members = {}
    for r in ov.iter_rows(min_row=6, values_only=True):
        if r[0] == 'TOTAL' or r[7] is None:
            continue
        members[r[7]] = dict(id=r[7], nick=r[8], role=r[6], superAgent=dash_(r[2]), agent=dash_(r[4]),
                             games=int(num(r[9])), hands=int(num(r[10])), fee=num(r[11]), bbj=num(r[15]), pnl=num(r[19]),
                             pnlNLH=0, pnlPLO5=0, pnlMTT=0, rakeNLH=0, rakePLO5=0, chips=0, credits=0)

    # Member Statistics: P&L and rake split by game. cols: 9 ring NLH P&L, 11 ring PLO5 P&L,
    # 28 MTT NLH P&L, 37 ring NLH rake, 39 ring PLO5 rake (see the sheet's 3 header rows).
    ms = wb['Member Statistics']
    for r in ms.iter_rows(min_row=7, values_only=True):
        if r[7] in members:
            members[r[7]].update(pnlNLH=num(r[9]), pnlPLO5=num(r[11]), pnlMTT=num(r[28]),
                                 rakeNLH=num(r[37]), rakePLO5=num(r[39]))

    # Club Member Balance: 9 chips, 10 credits.
    bal = wb['Club Member Balance']
    for r in bal.iter_rows(min_row=6, values_only=True):
        if r[7] in members:
            members[r[7]].update(chips=num(r[9]), credits=num(r[10]))
        elif r[7] and IDRE.match(str(r[7])):
            members[r[7]] = dict(id=r[7], nick=r[8], role=r[6], superAgent=dash_(r[2]), agent=dash_(r[4]),
                                 games=0, hands=0, fee=0, bbj=0, pnl=0, pnlNLH=0, pnlPLO5=0, pnlMTT=0,
                                 rakeNLH=0, rakePLO5=0, chips=num(r[9]), credits=num(r[10]))

    # Ring Game Detail: blocks of header lines, then player rows, then a Total row.
    rd = wb['Ring Game Detail']
    sessions, cur = [], None
    for r in rd.iter_rows(values_only=True):
        a = r[0]
        if isinstance(a, str) and a.startswith('Start/End Time'):
            m = re.search(r'(\d{4}-\d\d-\d\d \d\d:\d\d:\d\d) ~ (\d{4}-\d\d-\d\d \d\d:\d\d:\d\d)', a)
            cur = dict(start=m.group(1), end=m.group(2), players=[])
            sessions.append(cur)
        elif isinstance(a, str) and a.startswith('Table Name'):
            cur['table'] = re.search(r'Table Name : (.*?) , Creator', a).group(1).strip()
        elif isinstance(a, str) and a.startswith('Table Information'):
            cur['game'] = re.search(r'Game : (\S+)', a).group(1)
            cur['blinds'] = re.search(r'Blinds : (\S+)', a).group(1)
        elif isinstance(a, str) and IDRE.match(a) and cur is not None:
            # 2 buy-in, 3 cashout, 4 hands, 10 rake, 11 P&L
            cur['players'].append(dict(id=a, buyin=num(r[2]), cashout=num(r[3]), hands=int(num(r[4])),
                                       rake=num(r[10]), pnl=num(r[11])))
        elif a == 'Total' and cur is not None:
            cur.update(buyin=num(r[2]), hands=int(num(r[4])), bbj=num(r[8]), rake=num(r[10]))

    # MTT Statistics (one row per tournament) + MTT Detail (one row per player).
    mt = wb['MTT Statistics']
    mtts = []
    for r in mt.iter_rows(min_row=7, values_only=True):
        if r[0] and r[0] != 'TOTAL':
            mtts.append(dict(title=r[0], game=r[2], start=r[12], end=r[13], entries=int(num(r[15])),
                             gtd=num(r[9]), prize=num(r[24]), overlay=num(r[26]),
                             reentryBuyin=num(r[20]), reentryFee=num(r[22])))
    res = []
    for r in wb['MTT Detail'].iter_rows(values_only=True):
        if isinstance(r[0], str) and IDRE.match(r[0]):
            res.append(dict(id=r[0], hands=int(num(r[10])), prize=num(r[12]), winnings=num(r[14]),
                            reentry=num(r[6]) + num(r[8])))
    if mtts:
        mtts[0]['results'] = res  # note: assumes one tournament per week, as in the sample
    return dict(meta=meta, members=list(members.values()), sessions=sessions, mtts=mtts)


def to_datasets(d):
    nick = {m['id']: m['nick'] for m in d['members']}
    stake = lambda s: f"{s['game']} {s['blinds'].split('(')[0]}"
    sessions, ps = [], []
    for i, s in enumerate(sorted(d['sessions'], key=lambda x: x['start']), 1):
        a, b = datetime.fromisoformat(s['start']), datetime.fromisoformat(s['end'])
        sessions.append(dict(session=i, table=s['table'], game=s['game'], stake=stake(s), blinds=s['blinds'],
                             start=s['start'], end=s['end'], minutes=round((b - a).total_seconds() / 60, 1),
                             players=len(s['players']), hands=s['hands'], buyin=s['buyin'], rake=s['rake'], bbj=s['bbj']))
        for p in s['players']:
            ps.append(dict(session=i, start=s['start'], game=s['game'], stake=stake(s), id=p['id'],
                           player=nick.get(p['id'], p['id']), buyin=p['buyin'], cashout=p['cashout'],
                           hands=p['hands'], rake=p['rake'], pnl=p['pnl']))
    members = [dict(id=m['id'], player=m['nick'], role=m['role'], agent=m['agent'] or 'Direct',
                    superAgent=m['superAgent'] or '', games=m['games'], hands=m['hands'], fee=m['fee'], bbj=m['bbj'],
                    pnl=m['pnl'], pnlNLH=m['pnlNLH'], pnlPLO5=m['pnlPLO5'], pnlMTT=m['pnlMTT'],
                    chips=m['chips'], credits=m['credits']) for m in d['members']]
    mtt, res = [], []
    if d['mtts']:
        t = d['mtts'][0]
        mtt = [dict(title=t['title'], game=t['game'], start=t['start'], end=t['end'], entries=t['entries'],
                    gtd=t['gtd'], prizePaid=t['prize'], overlay=t['overlay'],
                    reentryBuyin=t['reentryBuyin'], reentryFee=t['reentryFee'])]
        res = [dict(id=r['id'], player=nick.get(r['id'], r['id']), hands=r['hands'], reentries=r['reentry'],
                    prize=r['prize'], net=r['winnings']) for r in sorted(t.get('results', []), key=lambda r: -r['prize'])]
    return dict(sessions=sessions, player_sessions=ps, members=members, mtt=mtt, mtt_results=res)


if __name__ == '__main__':
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(__file__), '..', 'data')
    os.makedirs(out, exist_ok=True)
    data = to_datasets(read_export(src))
    for name, rows in data.items():
        with open(os.path.join(out, f'{name}.json'), 'w', encoding='utf-8') as f:
            json.dump(rows, f, ensure_ascii=False, separators=(',', ':'))
        print(f'{name}: {len(rows)} rows')
    with open(os.path.join(out, 'data.js'), 'w', encoding='utf-8') as f:
        f.write('window.CK_DATA = ' + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print('wrote data.js')
